"""Сборка финального .docx возражения (python-docx).

Здесь же происходит ВОССТАНОВЛЕНИЕ ПДн из карты замен — локально, в самый
последний момент, уже после того как вся работа с LLM завершена. Реальные
ФИО/адреса появляются только в этом файле и никуда не отправляются.

В колонтитул и шапку добавляется дисклеймер «ПРОЕКТ. Требует проверки юристом».
"""

from __future__ import annotations

from tools.anonymizer import restore

DISCLAIMER = (
    "ПРОЕКТ. Подготовлен с помощью ИИ. НЕ является юридической консультацией. "
    "Требует обязательной проверки квалифицированным юристом перед подачей в суд."
)


def _r(text: str, pii_map: dict[str, str]) -> str:
    """Восстановить ПДн в строке."""
    return restore(text or "", pii_map or {})


def build_objection_docx(
    facts: dict,
    sections: dict[str, str],
    proshu_sud: list[str],
    prilozheniya: list[str],
    pii_map: dict[str, str],
    output_path: str,
) -> str:
    """Собрать .docx возражения и сохранить по output_path.

    Args:
        facts: реквизиты дела (court, case_number, plaintiff, defendant, ...).
        sections: {название_раздела: текст} — содержательные разделы.
        proshu_sud: список просьб для блока «ПРОШУ СУД».
        prilozheniya: список приложений.
        pii_map: карта замен токен->значение для восстановления ПДн.
        output_path: путь сохранения .docx.

    Returns:
        Путь к сохранённому файлу.
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    pii_map = pii_map or {}
    doc = Document()

    # --- Дисклеймер сверху ---
    disc = doc.add_paragraph()
    run = disc.add_run(DISCLAIMER)
    run.bold = True
    run.font.size = Pt(10)

    # --- Шапка: суд, стороны, дело ---
    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    head.add_run(
        f"В {_r(facts.get('court', '________'), pii_map)}\n"
        f"Ответчик: {_r(facts.get('defendant', '________'), pii_map)}\n"
        f"Истец: {_r(facts.get('plaintiff', '________'), pii_map)}\n"
        f"Дело № {_r(facts.get('case_number', '________'), pii_map)}"
    )

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = title.add_run("ВОЗРАЖЕНИЯ на исковое заявление")
    t.bold = True
    t.font.size = Pt(14)

    # --- Содержательные разделы ---
    for name, body in sections.items():
        h = doc.add_paragraph()
        h.add_run(_r(name, pii_map)).bold = True
        doc.add_paragraph(_r(body, pii_map))

    # --- ПРОШУ СУД ---
    if proshu_sud:
        h = doc.add_paragraph()
        h.add_run("ПРОШУ СУД:").bold = True
        for i, item in enumerate(proshu_sud, 1):
            doc.add_paragraph(f"{i}. {_r(item, pii_map)}")

    # --- Приложения ---
    if prilozheniya:
        h = doc.add_paragraph()
        h.add_run("Приложения:").bold = True
        for i, item in enumerate(prilozheniya, 1):
            doc.add_paragraph(f"{i}. {_r(item, pii_map)}")

    # --- Дисклеймер в колонтитуле каждой страницы ---
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fr = fp.add_run(DISCLAIMER)
    fr.italic = True
    fr.font.size = Pt(8)

    doc.save(output_path)
    return output_path
